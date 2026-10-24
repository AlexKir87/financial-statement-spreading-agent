"""Recalculate a copy of a workbook headlessly and scan it for formula errors.

Follows [[pattern-xlsx-recalc-and-formula-verification]]. XlsxWriter files carry the Python-side
value as the cached result, and a plain LibreOffice conversion keeps stale caches, so the copy is
first round-tripped through openpyxl (which drops cached values) and then converted, which forces
a calculation. Every formula cell is then compared with the cached value the writer stored.
The check fails closed: no LibreOffice, a timeout or an unreadable copy is a failed verification.
The original file is never modified.
"""

import math
import shutil
import subprocess  # noqa: S404 - fixed argument list, no shell
import tempfile
from pathlib import Path
from typing import Final

import openpyxl

from app.state import WorkbookVersion
from app.workbook.models import VerificationResult

FORMULA_ERRORS: Final = frozenset(
    {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}
)
CONVERT_TIMEOUT_SECONDS: Final = 120
NUMBER_TOLERANCE: Final = 1e-6


def _soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def _recalculate(source: Path, workdir: Path, soffice: str) -> Path:
    stripped = workdir / "in" / "copy.xlsx"
    stripped.parent.mkdir()
    openpyxl.load_workbook(source).save(stripped)  # formulas kept, cached values dropped
    out = workdir / "out"
    out.mkdir()
    subprocess.run(  # noqa: S603 - arguments are fixed paths, not user text
        [
            soffice,
            "--headless",
            "--norestore",
            f"-env:UserInstallation=file://{workdir}/profile",
            "--convert-to",
            "xlsx:Calc MS Excel 2007 XML",
            "--outdir",
            str(out),
            str(stripped),
        ],
        check=True,
        capture_output=True,
        timeout=CONVERT_TIMEOUT_SECONDS,
    )
    return out / "copy.xlsx"


def _same(cached: object, recalculated: object) -> bool:
    if isinstance(cached, bool) or isinstance(recalculated, bool):
        return cached == recalculated
    if isinstance(cached, int | float) and isinstance(recalculated, int | float):
        return math.isclose(cached, recalculated, rel_tol=1e-9, abs_tol=NUMBER_TOLERANCE)
    return cached == recalculated


def verify_workbook(path: str | Path, *, min_formulas: int = 1) -> VerificationResult:
    source = Path(path)
    soffice = _soffice()
    if soffice is None:
        return VerificationResult(
            ok=False, formulas_checked=0, problems=("libreoffice_unavailable",)
        )
    with tempfile.TemporaryDirectory(prefix="verify-") as tmp:
        try:
            return _scan(source, _recalculate(source, Path(tmp), soffice), min_formulas)
        except Exception as exc:  # noqa: BLE001 - fail closed on any unreadable or corrupt file
            return VerificationResult(
                ok=False,
                formulas_checked=0,
                problems=(f"recalculation_failed: {type(exc).__name__}",),
            )


def _scan(source: Path, recalculated: Path, min_formulas: int) -> VerificationResult:
    formulas = openpyxl.load_workbook(source)
    cached = openpyxl.load_workbook(source, data_only=True)
    fresh = openpyxl.load_workbook(recalculated, data_only=True)
    problems: list[str] = []
    checked = 0
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                checked += 1
                where = f"{ws.title}!{cell.coordinate}"
                new = fresh[ws.title][cell.coordinate].value
                if new is None or (isinstance(new, str) and new in FORMULA_ERRORS):
                    problems.append(
                        f"{where}: recalculated to {new if new is not None else 'nothing'}"
                    )
                elif not _same(cached[ws.title][cell.coordinate].value, new):
                    problems.append(f"{where}: cached value differs from the recalculated one")
    if checked < min_formulas:
        problems.append(f"expected at least {min_formulas} formulas, found {checked}")
    return VerificationResult(ok=not problems, formulas_checked=checked, problems=tuple(problems))


def verify_version(version: WorkbookVersion) -> tuple[WorkbookVersion, VerificationResult]:
    """Verify a written version and return it with `verified` set from the result."""
    result = verify_workbook(version.path)
    return version.model_copy(update={"verified": result.ok}), result
