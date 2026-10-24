"""Read analyst corrections back from a workbook's `Mapping` sheet (plan step 11).

The writer leaves four yellow columns next to each mapping row (`CORRECTION_COLS`). An analyst
fills them in Excel and saves; this reader turns each filled row into a `Correction`. The sheet
is analyst input, so nothing in it is trusted: cells are read as plain text (a cell that looks
like a formula is just text here and is never evaluated or written back as a formula), lengths
are capped, and a bad row becomes an open issue that names the row and the field, never the
cell's text (a rejected correction later names its id, which is analyst text but only ever
written to the workbook as a string). The table location is this one function (spec,
Flagged concern 9), so moving the corrections to a separate file later changes only this module.
"""

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Final

from openpyxl import load_workbook
from pydantic import ValidationError

from app.state import Correction, OpenIssue
from app.workbook.writer import CORRECTION_COLS, MAPPING_COLS

MAPPING_SHEET: Final = "Mapping"
SUMMARY_SHEET: Final = "Summary"
MAX_FILE_BYTES: Final = 20_000_000
MAX_ROWS: Final = 50_000
MAX_COLS: Final = 50
MAX_ID_CHARS: Final = 64
MAX_WHO_CHARS: Final = 80
MAX_WHY_CHARS: Final = 2_000

_KEY_COLS: Final = ("Period", "Raw line id")
_FIELD_BY_COL: Final = {
    "Correction id": "correction_id",
    "Corrected standard line id": "standard_line_id",
    "Who": "who",
    "Why": "why",
}
_CAPS: Final = {
    "correction_id": MAX_ID_CHARS,
    "who": MAX_WHO_CHARS,
    "why": MAX_WHY_CHARS,
    "period": 120,
    "raw_line_id": 200,
}


class CorrectionsError(ValueError):
    """The file cannot be read as a corrections source (wrong file, wrong run, wrong layout)."""


@dataclass(frozen=True)
class CorrectionsRead:
    corrections: tuple[Correction, ...]
    open_issues: tuple[OpenIssue, ...]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _run_id(rows: Sequence[Sequence[Any]]) -> str | None:
    for row in rows:
        if len(row) >= 2 and _text(row[0]) == "Run id":
            return _text(row[1])
    return None


def _issue(code: str, row_no: int, detail: str) -> OpenIssue:
    return OpenIssue(code=code, where=f"{MAPPING_SHEET} row {row_no}", detail=detail)


def _parse_row(row_no: int, cells: Mapping[str, str]) -> Correction | OpenIssue | None:
    filled = {c: cells[c] for c in _FIELD_BY_COL if cells[c]}
    if not filled:
        return None
    missing = [c for c in _FIELD_BY_COL if not cells[c]]
    if missing:
        return _issue(
            "correction_incomplete",
            row_no,
            "correction skipped; fill in: " + ", ".join(missing),
        )
    values = {
        "correction_id": cells["Correction id"],
        "standard_line_id": cells["Corrected standard line id"],
        "who": cells["Who"],
        "why": cells["Why"],
        "period": cells["Period"],
        "raw_line_id": cells["Raw line id"],
    }
    too_long = sorted(f for f, cap in _CAPS.items() if len(values[f]) > cap)
    if too_long:
        return _issue("correction_invalid", row_no, "too long: " + ", ".join(too_long))
    try:
        return Correction(**values)
    except ValidationError as exc:
        fields = sorted({str(err["loc"][0]) for err in exc.errors() if err["loc"]})
        return _issue("correction_invalid", row_no, "invalid: " + ", ".join(fields))


def read_corrections(path: Path | str, *, run_id: str) -> CorrectionsRead:
    """Corrections in the `Mapping` sheet of `path`, in row order.

    `run_id` must match the workbook's Summary sheet, so a file from another run is refused.
    A correction id that appears twice with the same content is kept once; with different
    content the later row is dropped and flagged.
    """
    file = Path(path)
    if file.suffix.lower() != ".xlsx" or not file.is_file():
        raise CorrectionsError("the corrections source must be an existing .xlsx file")
    if file.stat().st_size > MAX_FILE_BYTES:
        raise CorrectionsError("the corrections file is too large")
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for a bad file
        raise CorrectionsError(f"cannot open the workbook ({type(exc).__name__})") from None
    try:
        if SUMMARY_SHEET not in wb.sheetnames or MAPPING_SHEET not in wb.sheetnames:
            raise CorrectionsError("the workbook has no Summary or Mapping sheet")
        found = _run_id(list(wb[SUMMARY_SHEET].iter_rows(max_row=20, values_only=True)))
        if found != run_id:
            raise CorrectionsError("the workbook belongs to another run")
        rows = wb[MAPPING_SHEET].iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLS, values_only=True)
        header = [_text(c) for c in next(rows, ())]
        needed = [*_KEY_COLS, *CORRECTION_COLS]
        if any(h not in header for h in needed) or any(h not in header for h in MAPPING_COLS):
            raise CorrectionsError("the Mapping sheet lacks the mapping or correction columns")
        position = {name: header.index(name) for name in needed}
        corrections: list[Correction] = []
        issues: list[OpenIssue] = []
        seen: dict[str, Correction] = {}
        for row_no, raw in enumerate(rows, start=2):
            if row_no > MAX_ROWS:
                raise CorrectionsError("the Mapping sheet has too many rows")
            cells = {name: _text(raw[i]) if i < len(raw) else "" for name, i in position.items()}
            parsed = _parse_row(row_no, cells)
            if parsed is None:
                continue
            if isinstance(parsed, OpenIssue):
                issues.append(parsed)
                continue
            earlier = seen.get(parsed.correction_id)
            if earlier is None:
                seen[parsed.correction_id] = parsed
                corrections.append(parsed)
            elif earlier != parsed:
                issues.append(
                    _issue(
                        "correction_id_reused",
                        row_no,
                        "same correction id with different content; this row was skipped",
                    )
                )
        return CorrectionsRead(tuple(corrections), tuple(issues))
    finally:
        wb.close()
