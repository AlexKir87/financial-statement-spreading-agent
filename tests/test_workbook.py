import inspect
import shutil
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import xlsxwriter
from pydantic import ValidationError

import app.workbook.writer as writer_module
from app.state import (
    ChangeEntry,
    CheckResult,
    DigitizedPage,
    ExceptionNote,
    MappedLine,
    OpenIssue,
    RawLine,
    ScaleInfo,
    Term,
)
from app.workbook import (
    LayoutRow,
    WorkbookData,
    WorkbookError,
    verify_version,
    verify_workbook,
    write_workbook,
)
from app.workbook.templates import CellRef, render

NEEDS_SOFFICE = pytest.mark.skipif(
    shutil.which("soffice") is None, reason="LibreOffice is required to recalculate"
)
CLOCK = datetime(2026, 10, 21, 6, 0, tzinfo=UTC)
PERIODS = ("FY2022", "FY2023")
SHEETS = [
    "Summary",
    "IS",
    "BS",
    "CF",
    "Source_Trace",
    "Reconciliations",
    "Open_Issues",
    "Change_Log",
    "Mapping",
]

# line id -> (statement, FY2022, FY2023); the FY2023 figures are the spec's base filing.
LINES = {
    "is_revenue": ("IS", 1800, 2000),
    "is_expenses": ("IS", 1400, 1500),
    "bs_cash": ("BS", 600, 900),
    "bs_other_assets": ("BS", 2300, 2600),
    "bs_total_liabilities": ("BS", 1100, 1200),
    "bs_retained_earnings": ("BS", 800, 1300),
    "bs_other_equity": ("BS", 1000, 1000),
    "cf_net_income": ("CF", 400, 500),
    "cf_net_change_in_cash": ("CF", 100, 300),
}
LAYOUT = (
    LayoutRow(line_id="is_revenue", label="Revenue"),
    LayoutRow(line_id="is_expenses", label="Expenses"),
    LayoutRow(
        line_id="is_net_income",
        label="Net income",
        template="difference",
        operands=("is_revenue", "is_expenses"),
    ),
    LayoutRow(line_id="bs_cash", label="Cash"),
    LayoutRow(line_id="bs_other_assets", label="Other assets"),
    LayoutRow(
        line_id="bs_total_assets",
        label="Total assets",
        template="sum",
        operands=("bs_cash", "bs_other_assets"),
    ),
    LayoutRow(line_id="bs_total_liabilities", label="Total liabilities"),
    LayoutRow(line_id="bs_retained_earnings", label="Retained earnings"),
    LayoutRow(line_id="bs_other_equity", label="Other equity"),
    LayoutRow(
        line_id="bs_total_liabilities_and_equity",
        label="Total liabilities and equity",
        template="sum",
        operands=("bs_total_liabilities", "bs_retained_earnings", "bs_other_equity"),
    ),
)


def _check(difference: int = 0, key: str = "bs_balances", period: str = "FY2023") -> CheckResult:
    return CheckResult(
        key=key,
        period=period,
        status="pass" if difference == 0 else "fail",
        expected=Decimal(3500),
        actual=Decimal(3500 + difference),
        difference=Decimal(difference),
        terms=(Term(line_id="bs_total_assets", period=period, value=Decimal(3500 + difference)),),
    )


def make_data(**overrides: object) -> WorkbookData:
    spread = {
        p: {lid: Decimal(vals[1 + i]) for lid, vals in LINES.items()} for i, p in enumerate(PERIODS)
    }
    raw, mapping = [], {}
    for p in PERIODS:
        for lid, (stmt, *_rest) in LINES.items():
            rid = f"{p}-{lid}"
            raw.append(
                RawLine(
                    raw_line_id=rid,
                    doc_id="d1",
                    page=5,
                    original_label=f"Label {lid}",
                    original_value_text="1,000",
                    statement=stmt,
                    period=p,
                    scale="units",
                    currency="USD",
                )
            )
            mapping[f"{p}|{rid}"] = MappedLine(
                raw_line_id=rid,
                period=p,
                standard_line_id=lid,
                confidence=0.99,
                rationale="rule hit",
                origin="rule",
            )
    fields: dict[str, object] = {
        "run_id": "run-0001",
        "company_ref": "company-a",
        "status": "delivered",
        "scale": ScaleInfo(unit="units", currency="USD"),
        "created": CLOCK,
        "spread": spread,
        "raw_lines": tuple(raw),
        "mapping": mapping,
        "layout": LAYOUT,
        "digitized": (
            DigitizedPage(doc_id="d1", page=5, digitizer_name="fake", digitizer_version="1.2"),
        ),
        "check_results": (
            _check(),
            _check(key="cf_cash_reconciles"),
            CheckResult(key="re_rollforward", period="FY2022", status="not_applicable"),
        ),
    }
    fields.update(overrides)
    return WorkbookData.model_validate(fields)


@pytest.fixture
def out(tmp_path: Path) -> Path:
    return tmp_path / "out"


# --- structure -----------------------------------------------------------------------------


def test_all_sheets_present_in_order(out: Path) -> None:
    v = write_workbook(make_data(), out)
    assert openpyxl.load_workbook(v.path).sheetnames == SHEETS


def test_file_is_named_from_run_id_and_version(out: Path) -> None:
    v = write_workbook(make_data(), out)
    assert Path(v.path).name == "run-0001_v1.xlsx"
    assert v.version == 1 and v.verified is None
    assert len(v.sha256) == 64


def test_statement_totals_are_live_formulas_not_numbers(out: Path) -> None:
    ws = openpyxl.load_workbook(write_workbook(make_data(), out).path)["BS"]
    totals = [c.value for r in ws.iter_rows() for c in r if c.data_type == "f"]
    assert "=SUM(C2,C3)" in totals and "=SUM(D5,D6,D7)" in totals
    assert openpyxl.load_workbook(out / "run-0001_v1.xlsx")["IS"]["D4"].value == "=D2-D3"


def test_every_number_has_a_trace_row(out: Path) -> None:
    ws = openpyxl.load_workbook(write_workbook(make_data(), out).path)["Source_Trace"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == len(LINES) * len(PERIODS)
    assert {"fake"} == {r[6] for r in rows} and {"1.2"} == {r[7] for r in rows}


def test_text_layer_pages_are_marked_as_such(out: Path) -> None:
    ws = openpyxl.load_workbook(write_workbook(make_data(digitized=()), out).path)["Source_Trace"]
    assert {r[6].value for r in ws.iter_rows(min_row=2)} == {"text layer"}


def test_number_without_source_trace_blocks_the_write(out: Path) -> None:
    data = make_data(raw_lines=())
    with pytest.raises(WorkbookError, match="without a source trace"):
        write_workbook(data, out)
    assert not list(out.glob("*.xlsx"))


def test_summary_exceptions_issues_and_change_log(out: Path) -> None:
    data = make_data(
        status="delivered_with_exceptions",
        exceptions=(
            ExceptionNote(
                check="bs_balances",
                period="FY2023",
                difference=Decimal(50),
                explanation="source does not foot",
                attempts_made=2,
            ),
        ),
        open_issues=(OpenIssue(code="scale_assumed", where="BS", detail="thousands inferred"),),
        change_log=(ChangeEntry(who="system", what="wrote v1", when=CLOCK, why="run finished"),),
    )
    wb = openpyxl.load_workbook(write_workbook(data, out).path)
    summary = [[c.value for c in r] for r in wb["Summary"].iter_rows()]
    assert ["Plug figures", 0] in [row[:2] for row in summary]
    assert any(row[0] == "bs_balances" and row[2] == 50 and row[3] == 2 for row in summary)
    assert [c.value for c in wb["Open_Issues"][2]] == ["scale_assumed", "BS", "thousands inferred"]
    assert wb["Change_Log"]["C2"].value == CLOCK.isoformat()


def test_mapping_sheet_has_editable_correction_columns(out: Path) -> None:
    ws = openpyxl.load_workbook(write_workbook(make_data(), out).path)["Mapping"]
    assert [c.value for c in ws[1]][7:] == [
        "Correction id",
        "Corrected standard line id",
        "Who",
        "Why",
    ]


# --- versions ------------------------------------------------------------------------------


def test_second_write_makes_v2_and_leaves_v1_byte_identical(out: Path) -> None:
    v1 = write_workbook(make_data(), out)
    before = Path(v1.path).read_bytes()
    changed = make_data(
        change_log=(ChangeEntry(who="analyst-7", what="remap", when=CLOCK, why="fix"),)
    )
    v2 = write_workbook(changed, out)
    assert (v2.version, Path(v2.path).name) == (2, "run-0001_v2.xlsx")
    assert Path(v1.path).read_bytes() == before
    assert v1.sha256 != v2.sha256


def test_identical_rewrite_returns_the_existing_version(out: Path) -> None:
    v1 = write_workbook(make_data(), out)
    again = write_workbook(make_data(), out)
    assert again == v1
    assert sorted(p.name for p in out.iterdir()) == ["run-0001_v1.xlsx"]


def test_a_write_never_replaces_an_existing_file(
    out: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    write_workbook(make_data(), out)
    squatter = out / "run-0001_v2.xlsx"
    real_existing = writer_module._existing_versions
    calls = {"n": 0}

    def stale_listing(output_dir: Path, run_id: str) -> dict[int, Path]:
        calls["n"] += 1  # first scan misses v2 (another writer is mid-flight), later ones see it
        found = real_existing(output_dir, run_id)
        return {k: v for k, v in found.items() if k != 2} if calls["n"] == 1 else found

    squatter.write_bytes(b"someone else's file")
    monkeypatch.setattr(writer_module, "_existing_versions", stale_listing)
    changed = make_data(change_log=(ChangeEntry(who="a", what="x", when=CLOCK, why="y"),))
    v = write_workbook(changed, out)
    assert v.version == 3
    assert squatter.read_bytes() == b"someone else's file"


def test_no_temporary_files_are_left_behind(out: Path) -> None:
    write_workbook(make_data(), out)
    assert [p.name for p in out.iterdir()] == ["run-0001_v1.xlsx"]


def test_run_id_cannot_escape_the_output_directory() -> None:
    with pytest.raises(ValidationError):
        make_data(run_id="../evil")


# --- the model cannot write formulas ---------------------------------------------------------


def test_layout_accepts_template_ids_only() -> None:
    with pytest.raises(ValidationError):
        LayoutRow(line_id="bs_total_assets", label="x", template="=SUM(A1)", operands=("bs_cash",))  # type: ignore[arg-type]
    with pytest.raises(ValidationError):
        LayoutRow(line_id="bs_total_assets", label="x", template="sum", operands=("=1+1",))


def test_no_public_writer_entry_takes_a_formula_string() -> None:
    params = inspect.signature(write_workbook).parameters
    assert list(params) == ["data", "output_dir"]
    assert not hasattr(writer_module, "write_formula_string")
    with pytest.raises(ValidationError):
        make_data(formula="=1+1")  # unknown fields are rejected, not written


def test_text_from_a_filing_is_written_as_a_string_never_a_formula(out: Path) -> None:
    hostile = '=HYPERLINK("http://x","y")'
    data = make_data()
    raw = tuple(r.model_copy(update={"original_label": hostile}) for r in data.raw_lines)
    data = data.model_copy(
        update={
            "raw_lines": raw,
            "open_issues": (
                OpenIssue(code="injection_text", where="p5", detail="+cmd|' /C calc'!A0"),
            ),
        }
    )
    wb = openpyxl.load_workbook(write_workbook(data, out).path)
    cell = wb["Source_Trace"]["E2"]
    assert cell.value == hostile and cell.data_type == "s"
    assert wb["Open_Issues"]["C2"].data_type == "s"


def test_templates_render_only_from_cell_references() -> None:
    a, b = CellRef(sheet="BS", row=1, col=2), CellRef(sheet="BS", row=2, col=2)
    assert render("sum", (a, b), on="BS") == "=SUM(C2,C3)"
    assert render("difference", (a, b), on="IS") == "=BS!C2-BS!C3"
    tol = CellRef(sheet="Reconciliations", row=0, col=1, abs_=True)
    assert (
        render("ok_check", (a, tol), on="BS") == '=IF(ABS(C2)<=Reconciliations!$B$1,"OK","CHECK")'
    )
    with pytest.raises(ValidationError):
        CellRef(sheet="Evil'!A1", row=0, col=0)  # type: ignore[arg-type]


def test_operand_on_another_statement_is_rejected(out: Path) -> None:
    bad = (LayoutRow(line_id="is_net_income", label="NI", template="sum", operands=("bs_cash",)),)
    with pytest.raises(WorkbookError, match="another statement"):
        write_workbook(make_data(layout=bad), out)


def test_missing_operand_value_is_rejected(out: Path) -> None:
    spread = {"FY2023": {"bs_cash": Decimal(1)}}
    layout = (
        LayoutRow(line_id="bs_cash", label="Cash"),
        LayoutRow(
            line_id="bs_total", label="Total", template="sum", operands=("bs_cash", "bs_other")
        ),
    )
    with pytest.raises(WorkbookError):
        write_workbook(make_data(spread=spread, layout=layout, raw_lines=(), mapping={}), out)


# --- recalculation and verification ----------------------------------------------------------


@NEEDS_SOFFICE
def test_recalculated_values_match_the_python_expected_values(out: Path) -> None:
    data = make_data(check_results=(_check(), _check(difference=50, key="cf_cash_reconciles")))
    v = write_workbook(data, out)
    result = verify_workbook(v.path)
    assert result.ok, result.problems
    assert result.formulas_checked > 10
    # Independent look at what a real spreadsheet engine computed, not at the writer's cache.
    recalculated = _recalc_values(Path(v.path), out.parent / "recalc")
    recon = recalculated["Reconciliations"]
    assert [recon[f"F{r}"] for r in (4, 5)] == [
        "OK",
        "CHECK",
    ]  # tolerance 1: a difference of 50 fails
    assert recon["E5"] == 50
    assert recalculated["BS"]["D4"] == 3500 and recalculated["IS"]["D4"] == 500
    summary = recalculated["Summary"]
    assert (summary["B10"], summary["B11"], summary["B12"]) == (1, 1, 0)


@NEEDS_SOFFICE
def test_difference_of_one_passes_and_two_fails_at_tolerance_one(out: Path) -> None:
    data = make_data(
        check_results=(
            _check(difference=1).model_copy(update={"status": "pass"}),
            _check(difference=2, key="cf_cash_reconciles"),
        )
    )
    v = write_workbook(data, out)
    assert verify_workbook(v.path).ok
    recon = _recalc_values(Path(v.path), out.parent / "recalc")["Reconciliations"]
    assert [recon["F4"], recon["F5"]] == ["OK", "CHECK"]


@NEEDS_SOFFICE
def test_injected_ref_error_blocks_verification(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    wb = xlsxwriter.Workbook(str(path))
    ws = wb.add_worksheet("BS")
    ws.write_number(0, 0, 1)
    ws.write_formula(1, 0, "=#REF!+A1", None, 1)
    wb.close()
    result = verify_workbook(path)
    assert not result.ok
    assert any("BS!A2" in p and "#REF!" in p for p in result.problems)


@NEEDS_SOFFICE
def test_wrong_cached_value_blocks_verification(tmp_path: Path) -> None:
    path = tmp_path / "stale.xlsx"
    wb = xlsxwriter.Workbook(str(path))
    ws = wb.add_worksheet("BS")
    ws.write_number(0, 0, 2)
    ws.write_formula(1, 0, "=A1*2", None, 0)  # the cache says 0, the formula says 4
    wb.close()
    result = verify_workbook(path)
    assert not result.ok and "cached value differs" in result.problems[0]


@NEEDS_SOFFICE
def test_verification_does_not_touch_the_original(out: Path) -> None:
    v = write_workbook(make_data(), out)
    before = Path(v.path).read_bytes()
    verified, result = verify_version(v)
    assert result.ok and verified.verified is True and verified.version == v.version
    assert Path(v.path).read_bytes() == before


def test_verification_fails_closed_without_libreoffice(
    out: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    v = write_workbook(make_data(), out)
    monkeypatch.setattr("shutil.which", lambda _name: None)
    verified, result = verify_version(v)
    assert not result.ok and result.problems == ("libreoffice_unavailable",)
    assert verified.verified is False


def _recalc_values(path: Path, work: Path) -> dict[str, dict[str, object]]:
    from app.workbook import verify as verify_module

    soffice = verify_module._soffice()
    assert soffice is not None
    work.mkdir(parents=True, exist_ok=True)
    recalculated = verify_module._recalculate(path, work, soffice)
    wb = openpyxl.load_workbook(recalculated, data_only=True)
    return {
        ws.title: {c.coordinate: c.value for r in ws.iter_rows() for c in r} for ws in wb.worksheets
    }


def test_formula_row_that_disagrees_with_its_traced_spread_value_blocks_the_write(
    out: Path,
) -> None:
    spread = {p: dict(v) for p, v in make_data().spread.items()}
    spread["FY2023"]["bs_total_assets"] = Decimal(1)
    raw = make_data().raw_lines + (
        RawLine(
            raw_line_id="x",
            doc_id="d1",
            page=5,
            original_label="Total assets",
            original_value_text="1",
            statement="BS",
            period="FY2023",
            scale="units",
            currency="USD",
        ),
    )
    mapping = {
        **make_data().mapping,
        "FY2023|x": MappedLine(
            raw_line_id="x",
            period="FY2023",
            standard_line_id="bs_total_assets",
            confidence=0.9,
            rationale="r",
            origin="rule",
        ),
    }
    with pytest.raises(WorkbookError, match="disagrees"):
        write_workbook(make_data(spread=spread, raw_lines=raw, mapping=mapping), out)


def test_unknown_statement_prefix_and_duplicate_layout_ids_are_rejected(out: Path) -> None:
    with pytest.raises(ValidationError):
        make_data(layout=(LAYOUT[0], LAYOUT[0]))
    with pytest.raises(WorkbookError):
        write_workbook(
            make_data(
                spread={"FY2023": {"xx_bad": Decimal(1)}}, layout=(), raw_lines=(), mapping={}
            ),
            out,
        )


@NEEDS_SOFFICE
def test_corrupt_file_fails_closed(tmp_path: Path) -> None:
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"junk")
    result = verify_workbook(bad)
    assert not result.ok and result.problems[0].startswith("recalculation_failed")


@NEEDS_SOFFICE
def test_workbook_with_formulas_flattened_to_numbers_is_not_verified(tmp_path: Path) -> None:
    path = tmp_path / "flat.xlsx"
    wb = xlsxwriter.Workbook(str(path))
    wb.add_worksheet("BS").write_number(0, 0, 999)
    wb.close()
    result = verify_workbook(path)
    assert not result.ok and "at least 1 formulas" in result.problems[0]
